"""Resume parser module for extracting raw text and structure from PDF and DOCX files."""

import io
import os
import re
from pathlib import Path
from typing import Union, BinaryIO

try:
    import pymupdf as fitz
except ImportError:
    import fitz  # PyMuPDF fallback
import docx


def _extract_page_text_blocks(page) -> str:
    """Extract text from a PDF page respecting column layout.

    Detects a two-column layout when text block x-positions form two clear groups.
    If detected, reads the left column completely before the right column (with any
    spanning header blocks read first and footer blocks read last).
    Otherwise, preserves the standard top-to-bottom reading order.
    """
    raw_blocks = page.get_text("blocks")
    # Filter text blocks (block_type == 0) with non-empty text
    text_blocks = [b for b in raw_blocks if b[6] == 0 and b[4].strip()]
    if len(text_blocks) < 2:
        return page.get_text("text")

    page_width = page.rect.width
    page_height = page.rect.height

    header_blocks = []
    footer_blocks = []
    cand_blocks = []

    for b in text_blocks:
        width = b[2] - b[0]
        is_spanning = width > 0.55 * page_width or (b[0] < 0.25 * page_width and b[2] > 0.75 * page_width)
        if is_spanning and b[1] < 0.25 * page_height:
            header_blocks.append(b)
        elif is_spanning and b[3] > 0.85 * page_height:
            footer_blocks.append(b)
        else:
            cand_blocks.append(b)

    if len(cand_blocks) < 2:
        return page.get_text("text")

    # Generate candidate vertical split points across the middle of the page
    cand_splits = [
        0.50 * page_width,
        0.33 * page_width,
        0.40 * page_width,
        0.60 * page_width,
        0.67 * page_width,
    ]
    for b1 in cand_blocks:
        for b2 in cand_blocks:
            if b1[2] < b2[0]:
                mid = (b1[2] + b2[0]) / 2
                if 0.20 * page_width <= mid <= 0.80 * page_width:
                    cand_splits.append(mid)

    best_split = None
    best_left, best_right = [], []
    max_sep = -1

    for split_x in set(cand_splits):
        left = [b for b in cand_blocks if b[2] <= split_x + 5 or (b[0] + b[2]) / 2 < split_x]
        right = [b for b in cand_blocks if b[0] >= split_x - 5 or (b[0] + b[2]) / 2 >= split_x]
        crossing = [b for b in cand_blocks if b[0] < split_x - 10 and b[2] > split_x + 10]

        if crossing or not left or not right:
            continue
        if len(left) + len(right) != len(cand_blocks):
            continue

        # Verify vertical overlap between the two candidate columns
        min_y_l, max_y_l = min(b[1] for b in left), max(b[3] for b in left)
        min_y_r, max_y_r = min(b[1] for b in right), max(b[3] for b in right)
        y_overlap = min(max_y_l, max_y_r) - max(min_y_l, min_y_r)
        if y_overlap <= 0:
            continue

        # Verify significant horizontal separation between left and right groups
        mean_l = sum(b[0] for b in left) / len(left)
        mean_r = sum(b[0] for b in right) / len(right)
        sep = mean_r - mean_l
        if sep < 0.15 * page_width:
            continue

        if sep > max_sep:
            max_sep = sep
            best_split = split_x
            best_left = left
            best_right = right

    if best_split is not None:
        ordered = (
            sorted(header_blocks, key=lambda b: (b[1], b[0]))
            + sorted(best_left, key=lambda b: (b[1], b[0]))
            + sorted(best_right, key=lambda b: (b[1], b[0]))
            + sorted(footer_blocks, key=lambda b: (b[1], b[0]))
        )
        return "\n".join(b[4].strip() for b in ordered if b[4].strip())

    return page.get_text("text")


def extract_text(file_path_or_bytes: Union[str, Path, bytes, bytearray, BinaryIO], filename: str = "") -> str:
    """Extract and clean text from a PDF or DOCX file.

    Parameters
    ----------
    file_path_or_bytes : str, Path, bytes, bytearray, or BinaryIO
        File system path to the document, or raw bytes / byte-like stream.
    filename : str, optional
        Original filename (e.g. 'resume.pdf'). Required if passing raw bytes
        or if the path does not contain a recognizable file extension.

    Returns
    -------
    str
        Cleaned, extracted plain text content.

    Raises
    ------
    FileNotFoundError
        If a file path was provided but does not exist.
    ValueError
        If the file format is unsupported, the input is empty,
        or no text could be extracted.
    """
    # Step 1: Identify filename and file extension
    target_name = filename
    if not target_name and isinstance(file_path_or_bytes, (str, Path)):
        target_name = str(file_path_or_bytes)

    if not target_name:
        raise ValueError("Filename with extension must be provided when passing raw bytes.")

    ext = Path(target_name).suffix.lower()

    # Step 2: Validate supported extensions
    if ext not in [".pdf", ".docx", ".txt"]:
        raise ValueError(
            f"Unsupported file format '{ext}'. Only .pdf, .docx, and .txt files are supported."
        )

    # Step 3: Check if file path exists or read bytes
    raw_bytes: bytes = b""
    is_path = isinstance(file_path_or_bytes, (str, Path))

    if is_path:
        path = Path(file_path_or_bytes)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")
        if path.stat().st_size == 0:
            raise ValueError("The provided file is empty (0 bytes).")
        raw_bytes = path.read_bytes()
    elif isinstance(file_path_or_bytes, (bytes, bytearray)):
        raw_bytes = bytes(file_path_or_bytes)
        if len(raw_bytes) == 0:
            raise ValueError("The provided byte stream is empty (0 bytes).")
    elif hasattr(file_path_or_bytes, "read"):
        raw_bytes = file_path_or_bytes.read()
        if len(raw_bytes) == 0:
            raise ValueError("The provided file-like object is empty (0 bytes).")
    else:
        raise ValueError("Unsupported input type for file_path_or_bytes.")

    extracted_chunks = []

    # Step 4: Extract text based on file format
    if ext == ".pdf":
        # Open PDF from byte stream using PyMuPDF (fitz)
        try:
            with fitz.open(stream=raw_bytes, filetype="pdf") as doc:
                if doc.page_count == 0:
                    raise ValueError("The PDF document contains no pages.")
                # Read text page by page
                for page_num in range(doc.page_count):
                    page = doc.load_page(page_num)
                    page_text = _extract_page_text_blocks(page)
                    if page_text:
                        extracted_chunks.append(page_text)
        except Exception as e:
            if isinstance(e, ValueError):
                raise
            raise ValueError(f"Failed to parse PDF document: {e}") from e

    elif ext == ".docx":
        # Open DOCX from byte stream using python-docx
        try:
            doc = docx.Document(io.BytesIO(raw_bytes))

            # 4a. Extract regular paragraph text
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    extracted_chunks.append(text)

            # 4b. Extract table text (rows and cells)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    # Deduplicate repeated adjacent cells (common in merged cells)
                    deduped_cells = []
                    for cell_text in row_cells:
                        if not deduped_cells or cell_text != deduped_cells[-1]:
                            deduped_cells.append(cell_text)
                    if deduped_cells:
                        extracted_chunks.append(" | ".join(deduped_cells))
        except Exception as e:
            if isinstance(e, ValueError):
                raise
            raise ValueError(f"Failed to parse DOCX document: {e}") from e

    elif ext == ".txt":
        # Decode plain text bytes as UTF-8 (with fallback for encoding issues)
        try:
            text_content = raw_bytes.decode("utf-8", errors="replace")
            for line in text_content.splitlines():
                stripped = line.strip()
                if stripped:
                    extracted_chunks.append(stripped)
        except Exception as e:
            raise ValueError(f"Failed to read TXT document: {e}") from e

    # Step 5: Clean and normalize the extracted text
    combined_text = "\n".join(extracted_chunks)

    # Normalize horizontal whitespace (tabs, multiple spaces)
    cleaned_lines = []
    for line in combined_text.splitlines():
        line = re.sub(r"[ \t]+", " ", line).strip()
        if line:
            cleaned_lines.append(line)

    result_text = "\n".join(cleaned_lines).strip()

    # Step 6: Verify extracted text is not empty
    if not result_text:
        raise ValueError("The document contains no extractable text.")

    return result_text


_EMAIL_PATTERN = re.compile(
    r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b"
)


def extract_email(text: str) -> str | None:
    """Extract primary contact email address from resume text.

    Parameters
    ----------
    text : str
        Cleaned or raw resume text.

    Returns
    -------
    str | None
        Normalized lowercase email address, or None if none found.
    """
    if not isinstance(text, str) or not text.strip():
        return None

    # First attempt: direct search with standard email pattern
    matches = _EMAIL_PATTERN.findall(text)
    if matches:
        for m in matches:
            cleaned = m.strip(".,;:()[]{}<>\"' ")
            if re.match(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$", cleaned):
                return cleaned.lower()

    # Second attempt: handle spaced @ signs from OCR/PDF extraction (e.g. "name @ domain.com")
    loose_matches = re.findall(
        r"\b([A-Za-z0-9._%+-]+)\s*@\s*([A-Za-z0-9.-]+\.[A-Za-z]{2,})\b",
        text,
    )
    if loose_matches:
        candidate = f"{loose_matches[0][0]}@{loose_matches[0][1]}".lower()
        if re.match(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$", candidate):
            return candidate

    return None


def infer_candidate_name(filename: str) -> str:
    """Infer candidate name from resume filename.

    Removes the words 'resume' and 'cv' (anywhere, case-insensitive),
    replaces underscores, hyphens, and digits-only tokens with spaces,
    collapses multiple spaces, and title-cases the result.
    """
    stem = Path(filename or "").stem
    name = stem.replace("_", " ").replace("-", " ")
    name = re.sub(r"\b(resume|cv)\b", " ", name, flags=re.IGNORECASE)
    name = re.sub(r"\b\d+\b", " ", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name.title() if name else stem


_SECTION_KEYWORDS_RE = re.compile(
    r"\b(resume|cv|profile|summary|education|skills)\b", re.IGNORECASE
)
_URL_PATTERN_RE = re.compile(
    r"(https?://|www\.|\.(?:com|org|net|io|edu|co|ai|dev|pk|app|tech)\b)", re.IGNORECASE
)
_SPACY_NLP = None


def _get_spacy_nlp():
    """Lazy load spaCy en_core_web_sm pipeline."""
    global _SPACY_NLP
    if _SPACY_NLP is None:
        import spacy
        try:
            _SPACY_NLP = spacy.load("en_core_web_sm")
        except OSError:
            _SPACY_NLP = None
    return _SPACY_NLP


def extract_candidate_name(text: str, filename: str = "") -> str:
    """Extract candidate name from resume text with fallback to filename.

    Strategy:
    (a) Check the first 5 non-empty lines for a line of 2-4 words in Title Case
        or ALL CAPS with no digits, '@', URLs or section keywords
        (Resume, CV, Profile, Summary, Education, Skills).
    (b) Otherwise take the first spaCy PERSON entity in the first 400 characters.
    (c) Otherwise fall back to the existing infer_candidate_name(filename).
    """
    # (a) Check first 5 non-empty lines
    if text and isinstance(text, str):
        non_empty_lines = [
            line.strip() for line in text.splitlines() if line.strip()
        ][:5]

        for line in non_empty_lines:
            # Must not contain digits, '@', URLs, or section keywords
            if re.search(r"\d", line):
                continue
            if "@" in line:
                continue
            if _URL_PATTERN_RE.search(line):
                continue
            if _SECTION_KEYWORDS_RE.search(line):
                continue

            words = line.split()
            if not (2 <= len(words) <= 4):
                continue

            # Must not contain punctuation symbols common in headers / job titles / bullets
            if re.search(r"[^\w\s\.\-']", line, re.UNICODE):
                continue

            is_caps = line.isupper()
            is_title = line.istitle() or all(
                w.istitle() or (w.isupper() and len(w) <= 2) for w in words
            )

            if is_title or is_caps:
                return line.strip()

        # (b) spaCy PERSON entity in first 400 characters
        sample_text = text[:400]
        nlp = _get_spacy_nlp()
        if nlp is not None:
            doc = nlp(sample_text)
            for ent in doc.ents:
                if ent.label_ == "PERSON":
                    candidate = ent.text.strip()
                    if candidate:
                        return candidate

    # (c) Fallback to infer_candidate_name(filename)
    return infer_candidate_name(filename)


_PHONE_PATTERN = re.compile(
    r"""(?x)
    (?<![\w\+])
    (?:
        # Pakistani +92 or 0092 (mobile: 3xx, 7 digits)
        (?:\+92|0092)[\s.-]?(?:\(0\))?[\s.-]?\(?3\d{2}\)?[\s.-]?\d{3}[\s.-]?\d{4}\b
        |
        # Pakistani 03xx
        (?:\(03\d{2}\)|03\d{2})[\s.-]?\d{3}[\s.-]?\d{4}\b
        |
        (?:\(03\d{2}\)|03\d{2})[\s.-]?\d{7}\b
        |
        # International formats starting with +
        \+[1-9]\d{0,2}[\s.-]?(?:\(\d{1,4}\)|\d{1,4})[\s.-]?(?:\d{1,4}[\s.-]?){1,3}\d{2,4}\b
    )
    """
)


def extract_phone(text: str) -> str | None:
    """Extract primary contact phone number from resume text.

    Supports Pakistani formats (+92 3xx xxxxxxx, 03xx-xxxxxxx) and common
    international formats (+1, +44, etc.). Returns the first valid number or None.
    Never returns digit runs that are years, IDs, or dates.
    """
    if not text or not isinstance(text, str):
        return None

    # Find candidate phone numbers
    for match in _PHONE_PATTERN.finditer(text):
        cand = match.group(0).strip()
        # Clean trailing punctuation
        cand = cand.rstrip(".,;:)>]'\"")

        digits = re.sub(r"\D", "", cand)
        # Phone numbers must have between 10 and 15 digits
        if 10 <= len(digits) <= 15:
            # Check that it's not a Pakistani CNIC (xxxxx-xxxxxxx-x has 13 digits but starts with 5 digits)
            if re.match(r"^\d{5}[-]?\d{7}[-]?\d$", cand):
                continue
            return cand

    return None


_URL_REGEX = re.compile(
    r"""(?xi)
    (?<![@\w])
    (?:
        https?://[^\s<>'"()]+
        |
        www\.[^\s<>'"()]+
        |
        \b(?:[a-zA-Z0-9-]+\.)+(?:com|org|net|io|dev|me|ai|co|app|tech|info|edu|gov|pk|uk|ca|de|fr|xyz)(?:/[^\s<>'"()]*)?
    )
    """
)


def extract_links(text: str) -> dict[str, str | None]:
    """Extract LinkedIn, GitHub, and other portfolio/personal links from text.

    Returns
    -------
    dict
        {"linkedin": str | None, "github": str | None, "other": str | None}
        with the first match of each. Handles links with and without https:// and www.
    """
    links: dict[str, str | None] = {"linkedin": None, "github": None, "other": None}
    if not text or not isinstance(text, str):
        return links

    for match in _URL_REGEX.finditer(text):
        url = match.group(0).strip()
        url = url.lstrip("(<[\"'").rstrip(".,;:)>]'\"")
        if not url or "@" in url:
            continue

        url_lower = url.lower()

        if "linkedin.com" in url_lower:
            if links["linkedin"] is None:
                links["linkedin"] = url
        elif "github.com" in url_lower:
            if links["github"] is None:
                links["github"] = url
        else:
            if links["other"] is None:
                links["other"] = url

    return links



