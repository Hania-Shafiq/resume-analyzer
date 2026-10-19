"""Comprehensive unit tests for the resume parser module (app/parser.py).

Covers:
- PDF extraction (valid text, multi-page, empty document)
- DOCX extraction (paragraphs, tables, deduplication of merged cells)
- TXT extraction (plain text, utf-8 decoding)
- Empty files (0-byte byte stream, 0-byte file path, 0-byte file-like object)
- Unsupported formats (.xyz, .png, .jpg, .csv, .exe, missing extension)
- Corrupted files (corrupted PDF bytes, corrupted DOCX bytes)
- Non-existent files and invalid input types
"""

import io
from pathlib import Path
import pytest
import docx

try:
    import pymupdf as fitz
except ImportError:
    import fitz

from app.parser import extract_text, extract_email


class TestParserFormats:
    """Tests for supported document formats."""

    def test_pdf_extraction_valid(self):
        """Ensure standard PDF text is extracted properly."""
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "Jane Doe - Data Scientist")
        page.insert_text((50, 80), "Machine Learning, FastAPI, PyTorch")
        pdf_bytes = doc.write()
        doc.close()

        extracted = extract_text(pdf_bytes, filename="sample.pdf")
        assert "Jane Doe - Data Scientist" in extracted
        assert "Machine Learning, FastAPI, PyTorch" in extracted

    def test_pdf_multipage_extraction(self):
        """Ensure multi-page PDF extracts text across all pages."""
        doc = fitz.open()
        p1 = doc.new_page()
        p1.insert_text((50, 50), "Page One: Senior Developer")
        p2 = doc.new_page()
        p2.insert_text((50, 50), "Page Two: Extensive Experience in Docker and Kubernetes")
        pdf_bytes = doc.write()
        doc.close()

        extracted = extract_text(pdf_bytes, filename="multipage.pdf")
        assert "Page One: Senior Developer" in extracted
        assert "Page Two: Extensive Experience in Docker and Kubernetes" in extracted

    def test_docx_extraction_paragraphs_and_tables(self):
        """Ensure DOCX paragraphs and tables are extracted accurately."""
        doc = docx.Document()
        doc.add_heading("Hania Shafiq - Resume", level=1)
        doc.add_paragraph("Experienced Software Engineer specializing in Python and NLP.")

        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Skill"
        table.cell(0, 1).text = "Proficiency"
        table.cell(1, 0).text = "Python"
        table.cell(1, 1).text = "Advanced"

        stream = io.BytesIO()
        doc.save(stream)
        docx_bytes = stream.getvalue()

        extracted = extract_text(docx_bytes, filename="sample.docx")
        assert "Hania Shafiq - Resume" in extracted
        assert "Experienced Software Engineer" in extracted
        assert "Python | Advanced" in extracted

    def test_docx_table_merged_cells_deduplication(self):
        """Ensure adjacent identical cells in tables (common in merged cells) are deduplicated."""
        doc = docx.Document()
        table = doc.add_table(rows=1, cols=3)
        # Simulate merged cells where adjacent cells have identical text
        table.cell(0, 0).text = "Project Alpha"
        table.cell(0, 1).text = "Project Alpha"
        table.cell(0, 2).text = "Lead Architect"

        stream = io.BytesIO()
        doc.save(stream)
        docx_bytes = stream.getvalue()

        extracted = extract_text(docx_bytes, filename="merged.docx")
        assert "Project Alpha | Lead Architect" in extracted

    def test_txt_extraction_valid(self):
        """Ensure UTF-8 text file is extracted properly."""
        txt_content = "Alex Murphy\nDevOps Engineer\nSkills: Docker, Kubernetes, Terraform\n5 years experience"
        extracted = extract_text(txt_content.encode("utf-8"), filename="resume.txt")
        assert "Alex Murphy" in extracted
        assert "Terraform" in extracted

    def test_txt_extraction_utf8_with_bom(self):
        """Ensure UTF-8 with BOM (Byte Order Mark) is parsed cleanly without BOM artifact."""
        txt_content = "\ufeffFatima Khan\nPython Developer\nSkills: Python, Django, PostgreSQL\n3 years experience"
        extracted = extract_text(txt_content.encode("utf-8-sig"), filename="resume_bom.txt")
        assert "Fatima Khan" in extracted
        assert not extracted.startswith("\ufeff")
        assert "PostgreSQL" in extracted

    def test_txt_extraction_latin1(self):
        """Ensure Latin-1 / non-UTF8 encoded text with special characters is handled gracefully."""
        txt_content = "René Müller\nSoftware Engineer\nSkills: C++, Qt, Python\nZürich, Switzerland"
        latin1_bytes = txt_content.encode("latin-1")
        extracted = extract_text(latin1_bytes, filename="resume_latin1.txt")
        assert "Müller" in extracted or "Ren" in extracted
        assert "Python" in extracted

    def test_txt_file_path_extraction(self, tmp_path):
        """Ensure extracting text from a .txt file on disk works correctly."""
        txt_path = tmp_path / "candidate_resume.txt"
        txt_path.write_text("Sarah Connor\nCybersecurity Specialist\nSkills: Linux, Python, Wireshark", encoding="utf-8")
        extracted = extract_text(txt_path)
        assert "Sarah Connor" in extracted
        assert "Cybersecurity Specialist" in extracted
        assert "Wireshark" in extracted

    def test_doc_extraction_valid(self):
        """Ensure .doc format is accepted and handled."""
        content = "John Doe\nBackend Developer\nSkills: Java, Spring Boot, MySQL"
        extracted = extract_text(content.encode("utf-8"), filename="resume.doc")
        assert "John Doe" in extracted
        assert "Spring Boot" in extracted

    def test_txt_resume_pipeline_integration(self):
        """Verify extracted text from TXT works seamlessly with all NLP pipeline components."""
        from app.skills import extract_skills, extract_education, extract_experience_years
        from app.parser import extract_email, extract_candidate_name, extract_phone
        from app.preprocess import light_clean

        txt_resume = (
            "David Kim\n"
            "david.kim@example.com | +1 555-234-5678\n"
            "Summary: Senior Cloud Engineer with 6 years of experience.\n"
            "Education: Bachelor of Science in Computer Science\n"
            "Skills: Python, FastAPI, Docker, Kubernetes, AWS, PostgreSQL\n"
            "Experience:\n"
            "- Led cloud infrastructure migration to AWS.\n"
        )
        extracted = extract_text(txt_resume.encode("utf-8"), filename="david_kim.txt")
        cleaned = light_clean(extracted)

        # 1. Candidate identity extraction
        name = extract_candidate_name(extracted, "david_kim.txt")
        assert "David Kim" in name
        email = extract_email(extracted)
        assert email == "david.kim@example.com"
        phone = extract_phone(extracted)
        assert phone is not None

        # 2. Skill extraction
        skills = extract_skills(cleaned)
        assert "python" in skills
        assert "docker" in skills
        assert "kubernetes" in skills
        assert "aws" in skills

        # 3. Education extraction
        edu = extract_education(cleaned)
        assert any(e.degree == "Bachelor's" for e in edu)

        # 4. Experience years extraction
        exp = extract_experience_years(cleaned)
        assert exp == 6


class TestParserEmptyInputs:
    """Tests for empty files and missing text."""

    def test_empty_bytes_input(self):
        """Ensure an empty byte stream raises a ValueError."""
        with pytest.raises(ValueError, match="empty"):
            extract_text(b"", filename="resume.pdf")

    def test_empty_bytearray_input(self):
        """Ensure an empty bytearray raises a ValueError."""
        with pytest.raises(ValueError, match="empty"):
            extract_text(bytearray(), filename="resume.docx")

    def test_empty_file_like_object(self):
        """Ensure an empty file-like stream raises a ValueError."""
        with pytest.raises(ValueError, match="empty"):
            extract_text(io.BytesIO(b""), filename="resume.pdf")

    def test_empty_file_on_disk(self, tmp_path):
        """Ensure an empty (0 byte) file path raises a ValueError."""
        empty_file = tmp_path / "empty_resume.pdf"
        empty_file.write_bytes(b"")

        with pytest.raises(ValueError, match="empty"):
            extract_text(empty_file)

    def test_pdf_blank_page_no_text(self):
        """Ensure a PDF with a blank page raises ValueError ('no extractable text')."""
        doc = fitz.open()
        doc.new_page()  # blank page
        pdf_bytes = doc.write()
        doc.close()

        with pytest.raises(ValueError, match="no extractable text"):
            extract_text(pdf_bytes, filename="blank_page.pdf")

    def test_document_with_whitespace_only(self, tmp_path):
        """Ensure a document containing only whitespace raises ValueError ('no extractable text')."""
        txt_file = tmp_path / "blank.txt"
        txt_file.write_text("   \n\t  \n   ", encoding="utf-8")

        with pytest.raises(ValueError, match="no extractable text"):
            extract_text(txt_file)


class TestParserUnsupportedFormats:
    """Tests for unsupported file formats and missing filenames."""

    @pytest.mark.parametrize("ext", [".xyz", ".png", ".jpg", ".csv", ".exe", ".json", ".zip"])
    def test_unsupported_file_extension(self, ext):
        """Ensure unsupported file extensions raise ValueError."""
        with pytest.raises(ValueError, match="Unsupported file format"):
            extract_text(b"some content", filename=f"resume{ext}")

    def test_missing_filename_on_raw_bytes(self):
        """Ensure passing raw bytes without filename raises ValueError."""
        with pytest.raises(ValueError, match="Filename with extension must be provided"):
            extract_text(b"some binary content")

    def test_unsupported_input_type(self):
        """Ensure passing unsupported types (e.g. int, dict) raises ValueError."""
        with pytest.raises(ValueError, match="Unsupported input type"):
            extract_text(12345, filename="resume.pdf")

    def test_nonexistent_file_path(self):
        """Ensure a non-existent file path raises FileNotFoundError."""
        with pytest.raises(FileNotFoundError):
            extract_text("non_existent_resume_file_12345.pdf")


class TestParserCorruptedFiles:
    """Tests for corrupted or malformed documents."""

    def test_corrupted_pdf_bytes(self):
        """Ensure corrupted PDF bytes raise ValueError with descriptive error message."""
        corrupted_bytes = b"%PDF-1.4\ncorrupted content that cannot be parsed as valid PDF"
        with pytest.raises(ValueError, match="Failed to parse PDF document"):
            extract_text(corrupted_bytes, filename="corrupted.pdf")

    def test_corrupted_docx_bytes(self):
        """Ensure corrupted DOCX bytes raise ValueError with descriptive error message."""
        corrupted_bytes = b"PK\x03\x04corrupted docx zip header not valid archive"
        with pytest.raises(ValueError, match="Failed to parse DOCX document"):
            extract_text(corrupted_bytes, filename="corrupted.docx")


class TestExtractEmail:
    """Tests for email extraction from resume text."""

    def test_standard_email(self):
        text = "Jane Doe\nEmail: jane.doe@example.com\nPhone: 123-456-7890"
        assert extract_email(text) == "jane.doe@example.com"

    def test_email_with_punctuation(self):
        text = "Contact: <candidate_123@subdomain.domain.org>."
        assert extract_email(text) == "candidate_123@subdomain.domain.org"

    def test_no_email_returns_none(self):
        text = "Just a resume with no email address listed."
        assert extract_email(text) is None


class TestTwoColumnPdfReadingOrder:
    """Tests for PDF two-column layout detection and reading order."""

    def test_single_column_pdf_reading_order(self):
        """Single-column PDF maintains standard top-to-bottom reading order."""
        doc = fitz.open()
        page = doc.new_page(width=600, height=800)
        page.insert_textbox(fitz.Rect(50, 40, 550, 80), "Header Title: Jane Doe")
        page.insert_textbox(fitz.Rect(50, 100, 550, 200), "Single Column Section 1: Summary")
        page.insert_textbox(fitz.Rect(50, 250, 550, 350), "Single Column Section 2: Experience")
        page.insert_textbox(fitz.Rect(50, 400, 550, 500), "Single Column Section 3: Education")
        pdf_bytes = doc.write()
        doc.close()

        extracted = extract_text(pdf_bytes, filename="single_col.pdf")
        lines = [line.strip() for line in extracted.splitlines() if line.strip()]

        assert lines == [
            "Header Title: Jane Doe",
            "Single Column Section 1: Summary",
            "Single Column Section 2: Experience",
            "Single Column Section 3: Education",
        ]

    def test_two_column_pdf_reads_left_column_completely_before_right(self):
        """Two-column PDF reads left column completely before right column."""
        doc = fitz.open()
        page = doc.new_page(width=600, height=800)

        # Spanning top header
        page.insert_textbox(fitz.Rect(50, 40, 550, 80), "Header: Alex Smith")

        # Interleaved column blocks: Left 1, Right 1, Left 2, Right 2
        page.insert_textbox(fitz.Rect(50, 100, 250, 250), "Left Col 1: Education\nBS in Computer Science")
        page.insert_textbox(fitz.Rect(350, 100, 550, 250), "Right Col 1: Experience\nSenior Software Engineer")
        page.insert_textbox(fitz.Rect(50, 300, 250, 500), "Left Col 2: Skills\nPython, Docker, FastAPI")
        page.insert_textbox(fitz.Rect(350, 300, 550, 500), "Right Col 2: Projects\nResume Analyzer Platform")

        pdf_bytes = doc.write()
        doc.close()

        extracted = extract_text(pdf_bytes, filename="two_col.pdf")

        pos_header = extracted.find("Header: Alex Smith")
        pos_left1 = extracted.find("Left Col 1: Education")
        pos_left2 = extracted.find("Left Col 2: Skills")
        pos_right1 = extracted.find("Right Col 1: Experience")
        pos_right2 = extracted.find("Right Col 2: Projects")

        assert pos_header != -1
        assert pos_left1 != -1
        assert pos_left2 != -1
        assert pos_right1 != -1
        assert pos_right2 != -1

        # Assert left column is read completely before right column
        assert pos_header < pos_left1 < pos_left2 < pos_right1 < pos_right2

