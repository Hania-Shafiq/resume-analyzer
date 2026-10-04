"""Tests for resume upload, full text extraction, and match analysis flow."""

import io
import unittest
from fastapi.testclient import TestClient
import docx

from app.main import app


class TestUploadAndMatchFlow(unittest.TestCase):
    """Test suite ensuring full resume text is preserved and used in matching."""

    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)

    def test_upload_returns_full_extracted_text(self):
        """Ensure /resume/upload returns both text_preview and complete extracted_text."""
        doc = docx.Document()
        doc.add_heading("Hania Shafiq - Senior Data Scientist", level=1)
        padding = (
            "Over 10 years of extensive experience leading engineering teams across global organizations, "
            "designing and architecting large-scale distributed cloud systems, and driving high-impact technical initiatives. "
            "Demonstrated success in architecting secure, reliable, and fault-tolerant infrastructure solutions. "
            "Proven track record of optimizing computational costs, establishing modern agile delivery workflows, "
            "and delivering complex digital transformation initiatives that serve millions of active end users daily. "
            "Regularly collaborated with cross-functional leadership including product managers, data scientists, "
            "and security officers to define technology roadmaps, mitigate risks, and elevate developer velocity."
        )
        doc.add_paragraph(padding)
        doc.add_paragraph("Technical competencies: Python, FastAPI, Docker, PostgreSQL, PyTorch, Kubernetes.")

        stream = io.BytesIO()
        doc.save(stream)
        docx_bytes = stream.getvalue()

        response = self.client.post(
            "/resume/upload",
            files={"file": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()

        # Check required fields exist
        self.assertIn("extracted_text", data)
        self.assertIn("text_preview", data)
        self.assertIn("char_count", data)

        # Extracted text must be full, not ending with ellipsis
        self.assertGreater(len(data["extracted_text"]), 800)
        self.assertFalse(data["extracted_text"].endswith("..."))
        self.assertIn("Kubernetes", data["extracted_text"])

        # Preview must be truncated to 500 chars + ellipsis
        self.assertTrue(data["text_preview"].endswith("..."))
        self.assertLessEqual(len(data["text_preview"]), 503)

    def test_matching_with_full_extracted_text(self):
        """Ensure matching with full resume text detects skills that would be missed if truncated."""
        prefix = (
            "Candidate Profile: John Doe\n"
            "Summary: Highly experienced software engineer and engineering manager with an extensive background "
            "of leading multi-disciplinary software teams, delivering complex cloud-native systems, optimizing "
            "high-throughput architectures, and establishing best-in-class software development lifecycles across "
            "multiple business units and continents. Championed engineering excellence, automated code review processes, "
            "and modern deployment standards.\n"
            "Work Experience: Directed engineering operations for enterprise platforms in financial technology and healthcare sectors. "
            "Supervised system stability, disaster recovery planning, and compliance certifications.\n"
        )
        skills_section = "Technical Skills and Tools: Kubernetes, PostgreSQL, FastAPI, AWS, Docker, Python."
        long_resume = prefix + skills_section
        self.assertGreater(len(long_resume), 600)
        self.assertGreater(long_resume.index("Kubernetes"), 500)

        jd = (
            "We are seeking a Backend Engineer.\n"
            "Required skills: Kubernetes, PostgreSQL, FastAPI, AWS.\n"
            "Preferred skills: Docker, Python."
        )

        # Match using full resume text
        resp_full = self.client.post(
            "/match",
            json={"resume_text": long_resume, "job_description": jd}
        )
        self.assertEqual(resp_full.status_code, 200)
        data_full = resp_full.json()

        # All required skills should be matched with the full resume
        self.assertIn("kubernetes", data_full["required_matched"])
        self.assertIn("postgresql", data_full["required_matched"])
        self.assertIn("fastapi", data_full["required_matched"])
        self.assertIn("aws", data_full["required_matched"])
        self.assertEqual(data_full["skill_overlap_score"], 1.0)

        # Contrast with truncated preview (which cuts off before the skills)
        truncated_preview = long_resume[:500] + "..."
        resp_trunc = self.client.post(
            "/match",
            json={"resume_text": truncated_preview, "job_description": jd}
        )
        self.assertEqual(resp_trunc.status_code, 200)
        data_trunc = resp_trunc.json()

        # Truncated preview fails to match the skills that occur past 500 chars
        self.assertNotIn("kubernetes", data_trunc["required_matched"])
        self.assertLess(data_trunc["skill_overlap_score"], data_full["skill_overlap_score"])


if __name__ == "__main__":
    unittest.main()
