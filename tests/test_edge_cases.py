"""Edge case test suite for Resume Analyzer.

Covers:
1. Resume with only 2 lines
2. Resume with tables (DOCX and PDF)
3. Resume with unusual characters and Urdu text
4. A very long Job Description (25,000+ chars)
5. A Job Description with no skills at all
"""

import io
import pytest
from fastapi.testclient import TestClient
import docx

try:
    import pymupdf as fitz
except ImportError:
    import fitz

from app.main import app
from app.parser import extract_text
from app.preprocess import light_clean
from app.skills import (
    extract_skills,
    extract_education,
    extract_experience_years,
    analyze_job_description,
)
from app.matcher import semantic_similarity, split_sentences
from app.recommender import recommend_roles
from app.ranker import rank_candidates


@pytest.fixture(scope="module")
def client():
    """Shared FastAPI TestClient fixture."""
    return TestClient(app)


# ===========================================================================
# Edge Case 1: Resume with Only 2 Lines
# ===========================================================================

class TestEdgeCaseTwoLineResume:
    """Tests for minimal resume containing only two lines."""

    TWO_LINE_RESUME = "Jane Doe\nPython Developer"

    def test_two_line_resume_parsing(self):
        """extract_text extracts the 2 lines cleanly."""
        extracted = extract_text(self.TWO_LINE_RESUME.encode("utf-8"), filename="resume.txt")
        lines = extracted.splitlines()
        assert len(lines) == 2
        assert lines[0] == "Jane Doe"
        assert lines[1] == "Python Developer"

    def test_two_line_resume_skill_extraction(self):
        """extract_skills extracts skills from 2 lines without error."""
        skills = extract_skills(self.TWO_LINE_RESUME)
        assert skills == ["python"]

    def test_two_line_resume_education_and_experience(self):
        """extract_education and extract_experience_years handle 2 lines safely."""
        edu = extract_education(self.TWO_LINE_RESUME)
        assert edu == []
        exp = extract_experience_years(self.TWO_LINE_RESUME)
        assert exp is None

    def test_two_line_resume_semantic_matching(self):
        """semantic_similarity runs safely on a 2-line resume."""
        jd = "Requirements: Python backend developer."
        score = semantic_similarity(self.TWO_LINE_RESUME, jd)
        assert isinstance(score, float)
        assert 0.0 <= score <= 1.0
        assert score > 0.30

    def test_two_line_resume_role_recommendation(self):
        """recommend_roles returns predictions for 2-line resume."""
        roles = recommend_roles(self.TWO_LINE_RESUME, top_k=3)
        assert len(roles) == 3
        assert any("Python" in r["role"] or "Developer" in r["role"] or "Software" in r["role"] for r in roles)

    def test_two_line_resume_api_match(self, client):
        """POST /match handles a 2-line resume."""
        resp = client.post(
            "/match",
            json={"resume_text": self.TWO_LINE_RESUME, "job_description": "Requirements: Python."},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["skill_overlap_score"] == 1.0
        assert "python" in data["required_matched"]


# ===========================================================================
# Edge Case 2: Resume with Tables
# ===========================================================================

class TestEdgeCaseResumeWithTables:
    """Tests for resumes with complex table layouts."""

    def test_docx_table_with_skills_and_experience(self):
        """Ensure DOCX tables with multi-column layout extract skills and experience."""
        doc = docx.Document()
        doc.add_heading("Candidate Experience Matrix", level=1)

        # Create skills table
        table = doc.add_table(rows=3, cols=3)
        table.cell(0, 0).text = "Category"
        table.cell(0, 1).text = "Skill"
        table.cell(0, 2).text = "Experience"

        table.cell(1, 0).text = "Backend"
        table.cell(1, 1).text = "FastAPI, Python, PostgreSQL"
        table.cell(1, 2).text = "5+ years of experience"

        table.cell(2, 0).text = "DevOps"
        table.cell(2, 1).text = "Docker, Kubernetes"
        table.cell(2, 2).text = "3 years"

        stream = io.BytesIO()
        doc.save(stream)
        docx_bytes = stream.getvalue()

        extracted = extract_text(docx_bytes, filename="matrix.docx")
        assert "FastAPI, Python, PostgreSQL" in extracted
        assert "5+ years of experience" in extracted

        skills = extract_skills(extracted)
        assert "python" in skills
        assert "fastapi" in skills
        assert "postgresql" in skills
        assert "docker" in skills
        assert "kubernetes" in skills

        exp_years = extract_experience_years(extracted)
        assert exp_years == 5

    def test_pdf_table_text_extraction(self):
        """Ensure PDF tabular layout preserves textual data."""
        doc = fitz.open()
        page = doc.new_page()
        # Insert table headers and rows as spaced coordinates
        page.insert_text((50, 50), "Skill           Level          Years")
        page.insert_text((50, 70), "Python          Expert         6 years of experience")
        page.insert_text((50, 90), "Docker          Advanced       4 years")
        pdf_bytes = doc.write()
        doc.close()

        extracted = extract_text(pdf_bytes, filename="table.pdf")
        assert "Python" in extracted
        assert "Docker" in extracted
        skills = extract_skills(extracted)
        assert "python" in skills
        assert "docker" in skills
        assert extract_experience_years(extracted) == 6


# ===========================================================================
# Edge Case 3: Resume with Unusual Characters and Urdu Text
# ===========================================================================

class TestEdgeCaseUnusualCharactersAndUrdu:
    """Tests for resumes with non-ASCII, Urdu script, emojis, and symbols."""

    URDU_RESUME = (
        "محمد علی - سافٹ ویئر انجینئر\n"
        "تجربہ: پائیتھون اور مشین لرننگ میں 5 سال کا پیشہ ورانہ تجربہ۔\n"
        "مہارتیں: Python, FastAPI, Docker, PostgreSQL, Machine Learning.\n"
        "تعلیم: Bachelor of Science in Computer Science.\n"
        "رابطہ: ali@example.com | فون: +92 300 1234567\n"
        "🚀 Passionate about building cloud native microservices! 💻 ✨"
    )

    def test_light_clean_preserves_urdu_and_normalizes_emojis(self):
        """light_clean should not crash or mangle Urdu script."""
        cleaned = light_clean(self.URDU_RESUME)
        assert "سافٹ ویئر انجینئر" in cleaned
        assert "پائیتھون" in cleaned
        assert "Python" in cleaned

    def test_extract_skills_from_mixed_urdu_english(self):
        """extract_skills extracts English technical skills embedded in Urdu resume."""
        skills = extract_skills(self.URDU_RESUME)
        assert "python" in skills
        assert "fastapi" in skills
        assert "docker" in skills
        assert "postgresql" in skills
        assert "machine learning" in skills

    def test_extract_education_from_mixed_text(self):
        """extract_education extracts degree from mixed script text."""
        edu = extract_education(self.URDU_RESUME)
        degrees = [e.degree for e in edu]
        assert "Bachelor's" in degrees

    def test_semantic_similarity_with_urdu_text(self):
        """semantic_similarity executes without UnicodeDecodeError or crash."""
        jd = "Requirements: Python, FastAPI, and Docker for backend microservices."
        score = semantic_similarity(self.URDU_RESUME, jd)
        assert isinstance(score, float)
        assert 0.0 <= score <= 1.0
        assert score > 0.40  # English skills and experience present

    def test_api_upload_with_urdu_text(self, client):
        """POST /resume/upload handles a resume containing Urdu text and emojis."""
        txt_bytes = self.URDU_RESUME.encode("utf-8")
        resp = client.post(
            "/resume/upload",
            files={"file": ("resume_urdu.txt", txt_bytes, "text/plain")},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert "سافٹ ویئر" in data["extracted_text"]
        assert "python" in data["skills"]
        assert "docker" in data["skills"]


# ===========================================================================
# Edge Case 4: A Very Long Job Description
# ===========================================================================

class TestEdgeCaseVeryLongJD:
    """Tests for an unusually long job description (> 20,000 characters)."""

    @classmethod
    def setup_class(cls):
        """Generate a massive JD text with thousands of words."""
        filler_para = (
            "We are an international technology conglomerate operating at global scale. "
            "Our engineering culture values collaboration, high performance, and innovation. "
            "Engineers work autonomously in cross-functional squads to solve mission-critical challenges.\n"
        )
        requirements_block = (
            "Requirements:\n"
            "- Strong proficiency in Python, FastAPI, and Go\n"
            "- Hands-on expertise with Docker, Kubernetes, and AWS\n"
            "- In-depth knowledge of PostgreSQL, Redis, and Kafka\n"
            "- 5+ years of software engineering experience\n"
            "- Bachelor's degree in Computer Science or related field\n"
        )
        preferred_block = (
            "Preferred:\n"
            "- Experience with Machine Learning and PyTorch\n"
            "- Knowledge of Terraform and CI/CD pipelines\n"
        )
        # Repeat filler paragraphs to build > 25,000 characters
        cls.LONG_JD = (filler_para * 70) + requirements_block + (filler_para * 30) + preferred_block

    def test_long_jd_length(self):
        """Verify the test JD is indeed very long."""
        assert len(self.LONG_JD) > 20000

    def test_long_jd_analysis(self):
        """analyze_job_description parses very long JD without crash or hanging."""
        analysis = analyze_job_description(self.LONG_JD)
        assert "python" in analysis.required_skills
        assert "fastapi" in analysis.required_skills
        assert "docker" in analysis.required_skills
        assert "machine learning" in analysis.preferred_skills
        assert analysis.min_experience_years == 5

    def test_long_jd_api_analyze(self, client):
        """POST /job/analyze handles very long JD payload."""
        resp = client.post("/job/analyze", json={"description": self.LONG_JD})
        assert resp.status_code == 200
        data = resp.json()
        assert "python" in data["required_skills"]
        assert "kubernetes" in data["required_skills"]

    def test_long_jd_matching(self):
        """semantic_similarity runs on very long JD."""
        resume = "Software Engineer with 5 years experience in Python, FastAPI, Docker, and Kubernetes."
        score = semantic_similarity(resume, self.LONG_JD)
        assert isinstance(score, float)
        assert 0.0 <= score <= 1.0


# ===========================================================================
# Edge Case 5: A Job Description with No Skills at All
# ===========================================================================

class TestEdgeCaseJDWithNoSkills:
    """Tests for job descriptions that list zero recognizable technical skills."""

    JD_NO_SKILLS = (
        "Warehouse Stocker\n\n"
        "Looking for a warehouse stocker to organize inventory boxes on shelves. "
        "Must be able to lift heavy boxes and load delivery vans each morning. "
        "No prior experience or specialized degrees needed. "
        "Walk in to our local distribution center on Oak Avenue to apply."
    )

    def test_jd_analysis_with_no_skills(self):
        """analyze_job_description returns empty lists when JD has no skills."""
        analysis = analyze_job_description(self.JD_NO_SKILLS)
        assert analysis.required_skills == []
        assert analysis.preferred_skills == []
        assert analysis.min_experience_years is None

    def test_match_endpoint_with_zero_skills_jd(self, client):
        """POST /match handles JD with no skills safely without ZeroDivisionError."""
        resume = "Senior Python Developer with 5 years experience in FastAPI and Docker."
        resp = client.post(
            "/match",
            json={"resume_text": resume, "job_description": self.JD_NO_SKILLS},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["total_required"] == 0
        assert data["total_preferred"] == 0
        assert data["required_matched"] == []
        assert data["preferred_matched"] == []
        # Score is 0.0 because there are no skills to overlap with
        assert data["skill_overlap_score"] == 0.0

    def test_skill_gap_with_zero_skills_jd(self, client):
        """GET /skill-gap with zero-skill JD returns empty gaps."""
        resp = client.get(
            "/skill-gap",
            params={
                "resume_text": "Python developer with Docker",
                "job_description": self.JD_NO_SKILLS,
            },
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["required_gap"] == []
        assert data["preferred_gap"] == []
        assert data["gap_count"] == 0

    def test_recommendations_with_zero_skills_jd(self, client):
        """GET /recommendations with zero-skill JD returns empty lists."""
        resp = client.get(
            "/recommendations",
            params={
                "resume_text": "Python developer with Docker",
                "job_description": self.JD_NO_SKILLS,
            },
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["recommended_to_learn"] == []
        assert data["nice_to_add"] == []

    def test_rank_candidates_when_all_scores_zero(self):
        """rank_candidates ranks candidates deterministically even when all match scores are 0."""
        candidates = [
            {
                "status": "analyzed",
                "filename": "candidate_a.pdf",
                "candidate_name": "Alice Smith",
                "match_score": 0.0,
                "skills_match_count": 0,
                "experience_years": 3,
                "professional_months": 36,
            },
            {
                "status": "analyzed",
                "filename": "candidate_b.pdf",
                "candidate_name": "Bob Jones",
                "match_score": 0.0,
                "skills_match_count": 0,
                "experience_years": 5,
                "professional_months": 60,
            },
        ]
        ranked_res = rank_candidates(candidates)
        ranked = ranked_res["ranked"]
        # Bob Jones has 60 months, Alice has 36 -> Bob ranked #1
        assert ranked[0]["candidate_name"] == "Bob Jones"
        assert ranked[1]["candidate_name"] == "Alice Smith"
